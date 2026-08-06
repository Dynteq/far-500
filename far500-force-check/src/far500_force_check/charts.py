"""Native XLSX-grafieken voor het setup_analyse-tabblad.

Combineert een gestapelde AreaChart (groen/rood envelope-gebied) met een
LineChart voor de krachtlijn en een marker-only LineChart voor
overschrijdingen, via openpyxl's `chart_a += chart_b`-patroon (gedeelde
categorie-as, axId/crossAx automatisch gekoppeld). Gevalideerd met een losse
spike vóór implementatie — zie far500-force-check/README.md.
"""

from __future__ import annotations

from openpyxl.chart import AreaChart, LineChart, Reference
from openpyxl.chart.marker import Marker

from .model import Analysis

GREEN = "9BE38A"
RED = "F0A3A3"
BLACK = "1A1A1A"
MARKER_RED = "CC0000"
BLUE = "4A90D9"
GREY_DASH = "999999"


def _write_helper_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list]):
    for j, h in enumerate(headers):
        ws.cell(row=start_row, column=start_col + j, value=h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(row=start_row + 1 + i, column=start_col + j, value=v)
    last_col = start_col + len(headers) - 1
    last_row = start_row + len(rows)
    return start_row, start_row + 1, last_row, start_col, last_col


def build_force_chart(ws, analysis: Analysis, start_row: int, start_col: int, anchor_cell: str) -> int:
    """Schrijf de hulptabel + gecombineerde grafiek (envelope/restant/kracht/overschrijding).

    Retourneert de eerstvolgende vrije rij (voor een eventuele 2e hulptabel).
    """
    samples = analysis.recording.samples
    results = analysis.sample_results

    peak_force = max((r.force_abs_N for r in results), default=0.0)
    peak_envelope = max((r.envelope_N for r in results), default=0.0)
    y_max = max(1.1 * peak_force, 1.1 * peak_envelope, 1.0)

    headers = ["t_s", "envelope_N", "restant_N", "force_abs_N", "overschrijding_N"]
    rows = [
        [s.t_s, r.envelope_N, y_max - r.envelope_N, r.force_abs_N, r.force_abs_N if not r.force_ok else None]
        for s, r in zip(samples, results)
    ]
    header_row, first_data_row, last_data_row, first_col, _ = _write_helper_table(ws, start_row, start_col, headers, rows)

    area = AreaChart()
    area.grouping = "stacked"
    area.overlap = 100
    area.add_data(
        Reference(ws, min_col=first_col + 1, max_col=first_col + 2, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    area.set_categories(Reference(ws, min_col=first_col, min_row=first_data_row, max_row=last_data_row))
    area.series[0].graphicalProperties.solidFill = GREEN
    area.series[0].graphicalProperties.line.noFill = True
    area.series[1].graphicalProperties.solidFill = RED
    area.series[1].graphicalProperties.line.noFill = True

    force_line = LineChart()
    force_line.add_data(
        Reference(ws, min_col=first_col + 3, max_col=first_col + 3, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    force_line.series[0].graphicalProperties.line.solidFill = BLACK
    force_line.series[0].graphicalProperties.line.width = 15000
    force_line.series[0].marker = Marker(symbol="none")
    force_line.series[0].smooth = False

    marker_line = LineChart()
    marker_line.add_data(
        Reference(ws, min_col=first_col + 4, max_col=first_col + 4, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    marker_line.series[0].marker = Marker(symbol="circle", size=6)
    marker_line.series[0].marker.graphicalProperties.solidFill = MARKER_RED
    marker_line.series[0].graphicalProperties.line.noFill = True
    marker_line.series[0].smooth = False

    combo = area
    combo += force_line
    combo += marker_line
    combo.title = "Bedienkracht vs. toegestane envelope"
    combo.x_axis.title = "t (s)"
    combo.y_axis.title = "N"
    combo.y_axis.scaling.min = 0
    combo.y_axis.scaling.max = y_max
    combo.height = 9
    combo.width = 22
    ws.add_chart(combo, anchor_cell)
    return last_data_row + 2


def build_height_chart(ws, analysis: Analysis, start_row: int, start_col: int, anchor_cell: str) -> int:
    """Optionele 2e grafiek: handvathoogte vs tijd met 135/170cm-referentielijnen."""
    criteria = analysis.criteria
    samples = analysis.recording.samples

    headers = ["t_s", "handle_height_cm", "H_THRESH", "H_MAX"]
    rows = [[s.t_s, s.handle_height_cm, criteria.H_THRESH, criteria.H_MAX] for s in samples]
    header_row, first_data_row, last_data_row, first_col, _ = _write_helper_table(ws, start_row, start_col, headers, rows)

    chart = LineChart()
    chart.add_data(
        Reference(ws, min_col=first_col + 1, max_col=first_col + 3, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=first_col, min_row=first_data_row, max_row=last_data_row))
    chart.series[0].graphicalProperties.line.solidFill = BLUE
    chart.series[0].graphicalProperties.line.width = 15000
    chart.series[0].marker = Marker(symbol="none")
    chart.series[0].smooth = False
    for s in chart.series[1:]:
        s.graphicalProperties.line.solidFill = GREY_DASH
        s.graphicalProperties.line.dashStyle = "dash"
        s.marker = Marker(symbol="none")
        s.smooth = False
    chart.title = "Handvathoogte vs. tijd"
    chart.x_axis.title = "t (s)"
    chart.y_axis.title = "cm"
    chart.height = 9
    chart.width = 22
    ws.add_chart(chart, anchor_cell)
    return last_data_row + 2
